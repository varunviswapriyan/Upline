import argparse
import collections
import datetime as dt
import hashlib
import json
import math
import os
import sys
import zipfile

import numpy as np
import pandas as pd
from scipy.spatial import cKDTree
from sklearn.linear_model import PoissonRegressor
from sklearn.metrics import roc_auc_score

sys.path.insert(0, os.path.dirname(__file__))
from dbf import read_dbf
from shp import read_shp

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "..", "data")
SALT = "upline-hackohio-2026"
GRID = 0.004
K_ANON = 10

STORMS = {
    "ice_feb2025": ("2025-02-11", "2025-02-20"),
    "helene_sep2024": ("2024-09-25", "2024-10-03"),
}


def log(*a):
    print("[pipeline]", *a, flush=True)


def ensure_extracted(raw):
    ext = os.path.join(raw, "extracted")
    for name in ("Data1", "Data2", "matched_upstream"):
        d = os.path.join(ext, name)
        z = os.path.join(raw, "raw", f"{name}.zip")
        if not os.path.isdir(d) or not os.listdir(d):
            log("extracting", z)
            os.makedirs(d, exist_ok=True)
            with zipfile.ZipFile(z) as zf:
                for m in zf.namelist():
                    if m.startswith("__MACOSX") or m.endswith("/"):
                        continue
                    target = os.path.join(d, os.path.basename(m))
                    with zf.open(m) as src, open(target, "wb") as dst:
                        dst.write(src.read())
    return ext


def layer(ext, folder, name, cols=None):
    base = os.path.join(ext, folder, name)
    df, _, _ = read_dbf(base + ".dbf", cols=cols)
    _, geo = read_shp(base + ".shp")
    return df, geo


def merc_to_ll(x, y):
    lon = x / 6378137.0 * 180 / math.pi
    lat = (2 * math.atan(math.exp(y / 6378137.0)) - math.pi / 2) * 180 / math.pi
    return lon, lat


def hav_m(lon1, lat1, lon2, lat2):
    R = 6371000.0
    p1, p2 = np.radians(lat1), np.radians(lat2)
    dphi = p2 - p1
    dl = np.radians(lon2 - lon1)
    a = np.sin(dphi / 2) ** 2 + np.cos(p1) * np.cos(p2) * np.sin(dl / 2) ** 2
    return 2 * R * np.arcsin(np.sqrt(a))


def xy(lon, lat):
    lon = np.asarray(lon, float)
    lat = np.asarray(lat, float)
    return np.c_[(lon + 80.0) * 88_500.0, (lat - 37.2) * 111_000.0]


def circ_norm(c):
    c = str(c).strip()
    if c.upper().startswith("CIR"):
        return c[3:].lstrip("0") or "0"
    return c.lstrip("0") or "0"


def hid(premise, k=0):
    h = hashlib.sha256((SALT + str(k) + str(premise)).encode()).hexdigest().upper()
    alphabet = "23456789ABCDEFGHJKMNPQRSTUVWXYZ"
    n = int(h[:12], 16)
    s = ""
    for _ in range(7):
        s += alphabet[n % len(alphabet)]
        n //= len(alphabet)
    return "UPL-" + s


def labor_model(xlsx):
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws = wb["Labor time material related"]
    rows = [r for r in ws.iter_rows(values_only=True)]
    avg = {}
    cur = None
    for r in rows:
        if isinstance(r[0], str) and "Average Hours" in r[0]:
            cur = r[0].replace(" CU Average Hours", "").replace(" Average Hours", "").replace(" by Size", "").strip()
            continue
        if cur and isinstance(r[0], (int, float)) and isinstance(r[1], (int, float)) and r[2] is None:
            avg[cur] = (float(r[0]), float(r[1]))
        if cur and isinstance(r[0], (int, float)) and isinstance(r[2], (int, float)) and cur in ("Anchor",):
            avg[cur] = (float(r[1]), float(r[2]))
        if cur == "Pole" and r[0] == "Medium":
            avg["Pole"] = (float(r[2]), float(r[3]))
        if cur == "Transformer" and r[0] == 50:
            avg["Transformer"] = (float(r[2]), float(r[3]))
    cnc = wb["CNC"]
    conn = None
    for r in cnc.iter_rows(values_only=True):
        if r[4] == 95 or (isinstance(r[4], (int, float)) and isinstance(r[5], float)):
            conn = (float(r[5]), float(r[6]))
    avg["Connector"] = conn or (0.215, 0.045)
    ilm = float(wb["Indirect Labor Multiplier"]["B1"].value)
    adders = {}
    for r in wb["Labor only work"].iter_rows(values_only=True):
        if r[0] and isinstance(r[3], (int, float)):
            adders[r[1] if r[1] and len(str(r[1])) < 40 else r[0]] = float(r[3])

    def ir(k, n=1.0):
        a = avg.get(k, (0.0, 0.0))
        return n * (a[0] + a[1])

    assemblies = {
        "pole": {
            "label": "Broken pole replacement",
            "parts": [("Pole", 1), ("Crossarm", 1), ("Insulator", 3), ("Pin", 3), ("Guying", 1),
                      ("Grounding", 1), ("Dead-End", 2), ("Connector", 4)],
        },
        "primary_wire": {
            "label": "Downed primary conductor (per span)",
            "parts": [("Dead-End", 2), ("Connector", 6), ("Insulator", 3), ("Pin", 1)],
            "extra": ("Structural Realignment", 0.544),
        },
        "secondary_wire": {
            "label": "Downed secondary / service",
            "parts": [("Secondary Conductor Attachment", 2), ("Connector", 4), ("Service", 1)],
        },
        "fuse": {"label": "Blown/damaged fuse cutout", "parts": [("Fuse Cutout", 1), ("Connector", 2)]},
        "recloser": {"label": "Recloser repair/replace", "parts": [("Recloser", 1), ("Equipment Lead", 2)]},
        "switch": {"label": "Switch repair/replace", "parts": [("Switch", 1), ("Connector", 3)]},
        "transformer": {"label": "Transformer replacement",
                        "parts": [("Transformer", 1), ("Equipment Lead", 1), ("Fuse Cutout", 1),
                                  ("Grounding", 1), ("Bracket", 1)]},
        "other": {"label": "Other equipment", "parts": [("Equipment Lead", 2), ("Connector", 2)]},
    }
    out = {}
    for k, a in assemblies.items():
        direct = sum(ir(p, n) for p, n in a["parts"])
        if "extra" in a:
            direct += a["extra"][1]
        out[k] = {"label": a["label"], "direct_hours": round(direct, 3),
                  "crew_hours": round(direct * ilm, 3),
                  "parts": [{"cu": p, "qty": n, "hours": round(ir(p, n), 3)} for p, n in a["parts"]]}
    return {"assemblies": out, "ilm": ilm, "cu_avg": {k: [round(v[0], 4), round(v[1], 4)] for k, v in avg.items()},
            "adders": {"access_restricted": 0.5, "confined_backyard": 0.5, "remote_site": 1.312,
                       "specialized_handling": 0.5}}


DAMAGE_CLASS = {
    "Pole": "pole", "primary": "primary_wire", "secondary": "secondary_wire", "Fuse Bank": "fuse",
    "Recloser Bank": "recloser", "Switch": "switch", "Transformer": "transformer",
    "Capacitor Bank": "other", "Step Transformer": "other", "Light": None, "": "primary_wire",
}
LOCAL_CLASSES = {"secondary_wire", "transformer"}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--raw", required=True, help="path to the challenge repo's data/ folder")
    args = ap.parse_args()
    raw = os.path.abspath(args.raw)
    ext = ensure_extracted(raw)
    os.makedirs(OUT, exist_ok=True)

    log("loading customers")
    cust, _ = layer(ext, "Data1", "Customer_Clip",
                    cols=["premise_nu", "aep_xfmr_l", "circuit_id", "kind_of_se", "service_po", "service__1", "city"])
    cust["lon"] = pd.to_numeric(cust.service_po, errors="coerce")
    cust["lat"] = pd.to_numeric(cust.service__1, errors="coerce")
    cust["circ"] = cust.circuit_id.map(circ_norm)
    cust = cust.dropna(subset=["lon", "lat"])
    log("customers", len(cust))

    log("parsing topology")
    topo_path = os.path.join(ext, "matched_upstream", "matched_upstream.txt")
    prem_xf = {}
    xf_first = {}
    xf_loc = {}
    xf_steps = collections.defaultdict(int)
    parent_votes = collections.defaultdict(collections.Counter)
    dev_info = {}
    with open(topo_path) as f:
        for line in f:
            r = line.rstrip("\n").split("\t")
            nodes = [x.split("*") for x in r[1:]]
            xf = nodes[0][1]
            prem_xf[r[0]] = xf
            xf_loc[xf] = nodes[0][3]
            prot = []
            steps = 0
            for n in nodes[1:]:
                if n[0] == "eo_power_xfrmr_inst":
                    steps += 1
                    continue
                prot.append(n)
                if n[1] not in dev_info:
                    kind = {"Circuit Breaker": "breaker", "Fuse": "fuse", "Sectionalizer": "sectionalizer"}.get(
                        n[2], "recloser")
                    dev_info[n[1]] = {"kind": kind, "subtype": n[2], "loc": n[3]}
            xf_steps[xf] = max(xf_steps[xf], steps)
            xf_first[xf] = prot[0][1]
            for a, b in zip(prot[:-1], prot[1:]):
                parent_votes[a[1]][b[1]] += 1
    parent = {d: c.most_common(1)[0][0] for d, c in parent_votes.items()}
    for d in list(parent):
        seen = {d}
        p = parent.get(d)
        while p is not None:
            if p in seen:
                parent.pop(d, None)
                break
            seen.add(p)
            p = parent.get(p)
    devs = sorted(dev_info)
    di = {d: i for i, d in enumerate(devs)}
    log("devices", len(devs), "transformers", len(xf_first), "premises in topology", len(prem_xf))

    cust = cust[cust.premise_nu.isin(prem_xf)].copy()
    cust["xf"] = cust.premise_nu.map(prem_xf)
    cust["oh"] = (cust.kind_of_se == "Overhead").astype(int)
    cust["ug"] = (cust.kind_of_se == "Underground").astype(int)

    xft = cust.groupby("xf").agg(lon=("lon", "mean"), lat=("lat", "mean"), n=("premise_nu", "size"),
                                 oh=("oh", "sum"), circ=("circ", "first")).reset_index()
    xft["zone"] = xft.xf.map(xf_first)
    xft["zi"] = xft.zone.map(di)
    xft["steps"] = xft.xf.map(xf_steps).fillna(0).astype(int)
    xidx = {x: i for i, x in enumerate(xft.xf)}

    dev_ll = {}
    for folder, name in [("Data1", "Fuse_Bank_Clip"), ("Data2", "Recloser_Bank_Clip"), ("Data2", "Switch_Clip")]:
        df, geo = layer(ext, folder, name, cols=["location_n", "aep_circui"])
        for loc, g in zip(df.location_n, geo):
            if g:
                dev_ll.setdefault(loc, g)
    for d in devs:
        info = dev_info[d]
        info["ll"] = dev_ll.get(info["loc"])

    children = collections.defaultdict(list)
    for d, p in parent.items():
        children[p].append(d)
    direct_cust = collections.Counter()
    direct_oh = collections.Counter()
    direct_xf = collections.Counter()
    for z, n, oh in zip(xft.zone, xft.n, xft.oh):
        direct_cust[z] += n
        direct_oh[z] += oh
        direct_xf[z] += 1
    down = {}

    def total(d):
        if d in down:
            return down[d]
        stack = [(d, False)]
        while stack:
            node, done = stack.pop()
            if done:
                down[node] = direct_cust[node] + sum(down[c] for c in children[node])
            elif node not in down:
                stack.append((node, True))
                for c in children[node]:
                    if c not in down:
                        stack.append((c, False))
        return down[d]

    for d in devs:
        total(d)
    depth = {}
    for d in devs:
        k, p = 0, parent.get(d)
        while p is not None and k < 30:
            k += 1
            p = parent.get(p)
        depth[d] = k

    log("snapping assets to transformers / zones")
    xt_xy = xy(xft.lon, xft.lat)
    tree_all = cKDTree(xt_xy)
    trees = {}
    for c, g in xft.groupby("circ"):
        trees[c] = (cKDTree(xy(g.lon, g.lat)), g.index.values)

    def nearest_xf(lon, lat, circ=None):
        p = xy([lon], [lat])[0]
        if circ is not None and circ in trees:
            t, idx = trees[circ]
            dist, j = t.query(p)
            return idx[j], dist
        dist, j = tree_all.query(p)
        return j, dist

    zf = collections.defaultdict(lambda: collections.Counter())
    backbone = []
    for folder, name, tag in [("Data1", "Primary_ExportFeatures", "pri"), ("Data2", "Secondary_ExportFeatures", "sec")]:
        df, geo = layer(ext, folder, name, cols=["aep_circui", "mounting", "report_len", "usage"])
        L = pd.to_numeric(df.report_len, errors="coerce").fillna(0).values
        mids = []
        for g in geo:
            if not g:
                mids.append((np.nan, np.nan))
                continue
            pts = g[0]
            mids.append(pts[len(pts) // 2])
        mids = np.array(mids)
        ok = ~np.isnan(mids[:, 0])
        circs = df.aep_circui.map(circ_norm).values
        for c in np.unique(circs):
            m = ok & (circs == c)
            if not m.any():
                continue
            if c in trees:
                t, idx = trees[c]
                _, j = t.query(xy(mids[m, 0], mids[m, 1]))
                xi = idx[j]
            else:
                _, xi = tree_all.query(xy(mids[m, 0], mids[m, 1]))
            zones = xft.zone.values[xi]
            mounts = df.mounting.values[m]
            for z, mt, l in zip(zones, mounts, L[m]):
                key = f"{tag}_{'oh' if mt == 'Overhead' else 'ug'}_ft"
                zf[z][key] += l
        if tag == "pri":
            for g, u, mt in zip(geo, df.usage, df.mounting):
                if g and u == "Backbone" and mt == "Overhead":
                    pts = g[0]
                    step = max(1, len(pts) // 4)
                    sp = pts[::step] + ([pts[-1]] if (len(pts) - 1) % step else [])
                    backbone.append([[round(a, 4), round(b, 4)] for a, b in sp])

    pdf, pgeo = layer(ext, "Data1", "Pole_Clip", cols=["material", "date_insta", "usage"])
    pll = np.array([g if g else (np.nan, np.nan) for g in pgeo])
    ok = ~np.isnan(pll[:, 0])
    _, pj = tree_all.query(xy(pll[ok, 0], pll[ok, 1]))
    yrs = pd.to_numeric(pdf.date_insta.str[:4], errors="coerce").values[ok]
    wood = pdf.material.isin(["Pine", "Wood", "Cedar", "Ponderosa Pine"]).values[ok]
    for z, y, w in zip(xft.zone.values[pj], yrs, wood):
        c = zf[z]
        c["poles"] += 1
        c["wood"] += int(w)
        if not np.isnan(y) and 1900 < y <= 2026:
            c["age_sum"] += 2026 - y
            c["age_n"] += 1
            c["old"] += int(y < 1975)

    fac = []
    for name, kind in [("Hospitals_Medic_Pairwis_Clip", "hospital"), ("Fire_Stations_E_Pairwis_Clip", "fire"),
                       ("Ambulance_Servi_Pairwis_Clip", "ems")]:
        df, geo = layer(ext, "Data1", name, cols=["NAME"])
        for g in geo:
            if not g:
                continue
            lon, lat = merc_to_ll(*g)
            j, dist = nearest_xf(lon, lat)
            if dist < 1500:
                fac.append({"kind": kind, "zone": xft.zone.values[j], "lon": lon, "lat": lat})
    crit_zone = collections.Counter(f["zone"] for f in fac)
    log("critical facilities matched", len(fac))

    log("snapping damage records")
    dmg, dgeo = layer(ext, "Data1", "arcos_prod_damage_09_28_Clip")
    dmg["lon"] = pd.to_numeric(dmg.Longitude, errors="coerce")
    dmg["lat"] = pd.to_numeric(dmg.Latitude, errors="coerce")
    dmg["rt"] = pd.to_datetime(dmg.Report_Tim, errors="coerce")
    dmg["rd"] = pd.to_datetime(dmg.Repaired_D, errors="coerce")
    dmg.loc[dmg.Repaired != "Yes", "rd"] = pd.NaT
    dmg["cls"] = dmg.Asset_Type.map(DAMAGE_CLASS)
    dmg["qty"] = pd.to_numeric(dmg.Media, errors="coerce").fillna(1).clip(lower=1)
    dmg["access"] = (dmg.Access_Iss == "Yes").astype(int)
    dmg["circ"] = dmg.Circuit.map(circ_norm)
    dmg = dmg.dropna(subset=["lon", "lat", "cls"]).copy()
    dmg["storm"] = "other"
    for s, (a, b) in STORMS.items():
        dmg.loc[(dmg.rt >= a) & (dmg.rt <= b), "storm"] = s
    xi_list, dist_list = [], []
    for lon, lat, c in zip(dmg.lon, dmg.lat, dmg.circ):
        j, d = nearest_xf(lon, lat, c)
        xi_list.append(j)
        dist_list.append(d)
    dmg["xi"] = xi_list
    dmg["snap_m"] = dist_list
    dmg["zone"] = xft.zone.values[dmg.xi.values]
    dev_by_loc_kind = collections.defaultdict(list)
    for d in devs:
        info = dev_info[d]
        if info["ll"]:
            dev_by_loc_kind[info["kind"]].append((d, info["ll"]))
    for kind, cls in [("fuse", "fuse"), ("recloser", "recloser")]:
        lst = dev_by_loc_kind[kind]
        if not lst:
            continue
        t = cKDTree(xy([p[1][0] for p in lst], [p[1][1] for p in lst]))
        m = dmg.cls == cls
        dd, jj = t.query(xy(dmg.lon[m], dmg.lat[m]))
        newz = [lst[j][0] if dist < 250 else z for j, dist, z in zip(jj, dd, dmg.zone[m])]
        dmg.loc[m, "zone"] = newz
    dmg["local"] = dmg.cls.isin(LOCAL_CLASSES)
    log("damage snapped; median snap distance m =", round(float(dmg.snap_m.median()), 1),
        "| by storm:", dmg.storm.value_counts().to_dict())

    Z = pd.DataFrame({"dev": devs})
    Z["kind"] = Z.dev.map(lambda d: dev_info[d]["kind"])
    Z["direct_cust"] = Z.dev.map(direct_cust).fillna(0)
    Z["direct_xf"] = Z.dev.map(direct_xf).fillna(0)
    Z["down_cust"] = Z.dev.map(down).fillna(0)
    Z["oh_frac"] = (Z.dev.map(direct_oh).fillna(0) / Z.direct_cust.clip(lower=1))
    for k in ["pri_oh_ft", "pri_ug_ft", "sec_oh_ft", "sec_ug_ft", "poles", "wood", "age_sum", "age_n", "old"]:
        Z[k] = Z.dev.map(lambda d: zf[d][k] if d in zf else 0).astype(float)
    Z["pole_age"] = np.where(Z.age_n > 0, Z.age_sum / Z.age_n.clip(lower=1), np.nan)
    Z["pole_age"] = Z.pole_age.fillna(Z.pole_age.median())
    Z["old_frac"] = Z.old / Z.age_n.clip(lower=1)
    Z["wood_frac"] = Z.wood / Z.poles.clip(lower=1)
    Z["depth"] = Z.dev.map(depth)
    Z["crit"] = Z.dev.map(crit_zone).fillna(0)
    zc = xft.groupby("zone").agg(lon=("lon", "mean"), lat=("lat", "mean"))
    Z = Z.join(zc, on="dev")
    for d in devs:
        pass
    Z["lon"] = Z.lon.fillna(Z.dev.map(lambda d: (dev_info[d]["ll"] or (np.nan, np.nan))[0]))
    Z["lat"] = Z.lat.fillna(Z.dev.map(lambda d: (dev_info[d]["ll"] or (np.nan, np.nan))[1]))
    cxy = xy(cust.lon, cust.lat)
    ctree = cKDTree(cxy)
    okz = ~Z.lon.isna()
    dens = np.zeros(len(Z))
    dens[okz.values] = [len(v) for v in ctree.query_ball_point(xy(Z.lon[okz], Z.lat[okz]), r=1000)]
    Z["density_1km"] = dens

    zone_dmg = dmg[~dmg.local]
    for s in list(STORMS) + ["other"]:
        cnt = zone_dmg[zone_dmg.storm == s].groupby("zone").size()
        Z[f"y_{s}"] = Z.dev.map(cnt).fillna(0)

    FEATS = {
        "log_pri_oh_mi": lambda z: np.log1p(z.pri_oh_ft / 5280),
        "log_pri_ug_mi": lambda z: np.log1p(z.pri_ug_ft / 5280),
        "log_poles": lambda z: np.log1p(z.poles),
        "oh_service_frac": lambda z: z.oh_frac,
        "wood_pole_frac": lambda z: z.wood_frac,
        "rural_index": lambda z: -np.log1p(z.density_1km) / 3.0,
        "log_downstream_cust": lambda z: np.log1p(z.down_cust),
        "is_recloser_zone": lambda z: (z.kind.isin(["recloser", "sectionalizer"])).astype(float),
        "is_breaker_zone": lambda z: (z.kind == "breaker").astype(float),
    }
    FRIENDLY = {
        "log_pri_oh_mi": "Overhead primary line in this section",
        "log_pri_ug_mi": "Underground primary line in this section",
        "log_sec_oh_mi": "Overhead neighborhood (secondary) wires",
        "log_poles": "Number of poles in this section",
        "oh_service_frac": "Homes served by overhead drops",
        "wood_pole_frac": "Share of wooden poles",
        "pole_age_yrs": "Average pole age",
        "pre1975_pole_frac": "Share of poles installed before 1975",
        "rural_index": "Rural / tree-exposed area",
        "is_recloser_zone": "Section protected by a recloser",
        "is_breaker_zone": "Main feeder section (substation breaker)",
        "log_downstream_cust": "Size of the network this section feeds",
    }
    X = pd.DataFrame({k: f(Z) for k, f in FEATS.items()}).fillna(0)
    mu_x = X.mean()
    sd_x = X.std().replace(0, 1)
    Xs = (X - mu_x) / sd_x

    log("fitting damage-rate model + back-test")

    def fit(train_storms):
        Xt = pd.concat([Xs] * len(train_storms), ignore_index=True)
        yt = np.concatenate([Z[f"y_{s}"].values for s in train_storms])
        for s in train_storms:
            Xt[f"storm_{s}"] = np.concatenate([np.full(len(Z), 1.0 if s == t else 0.0) for t in train_storms])
        m = PoissonRegressor(alpha=0.02, max_iter=3000, fit_intercept=False)
        m.fit(Xt.values, yt)
        coef = dict(zip(Xt.columns, m.coef_))
        return m, coef

    def eta(coef):
        return sum(Xs[k].values * coef[k] for k in FEATS)

    def severity_for_total(base_eta, total):
        return math.log(max(total, 1) / np.exp(base_eta).sum())

    backtests = []
    for train, test in [("ice_feb2025", "helene_sep2024"), ("helene_sep2024", "ice_feb2025")]:
        _, coef = fit([train])
        e = eta(coef)
        y = Z[f"y_{test}"].values
        a = severity_for_total(e, y.sum())
        lam = np.exp(e + a)
        p = 1 - np.exp(-lam)
        lab = (y > 0).astype(int)
        auc = roc_auc_score(lab, p)
        base = Z.direct_cust.values + 1e-6
        auc_base = roc_auc_score(lab, base)
        order = np.argsort(p)
        bins = np.array_split(order, 10)
        calib = [{"pred": float(p[b].mean()), "obs": float(lab[b].mean()), "n": int(len(b))} for b in bins]
        top = order[-int(len(order) * 0.2):]
        capture = float(lab[top].sum() / max(lab.sum(), 1))
        backtests.append({"train": train, "test": test, "auc": round(float(auc), 3),
                          "auc_baseline_customers_only": round(float(auc_base), 3),
                          "top20_capture": round(capture, 3), "damaged_zones": int(lab.sum()),
                          "zones": int(len(lab)), "calibration": calib})
        log(f"  back-test train={train} test={test}: AUC={auc:.3f} (baseline {auc_base:.3f}) "
            f"top-20% captures {capture:.0%} of damaged zones")

    _, coef = fit(["ice_feb2025", "helene_sep2024"])
    e_final = eta(coef)
    a_ice = coef["storm_ice_feb2025"]
    a_hel = coef["storm_helene_sep2024"]
    rng = np.random.default_rng(7)
    boot_eta = []
    n = len(Z)
    for b in range(24):
        idx = rng.integers(0, n, n)
        Xb = pd.concat([Xs.iloc[idx]] * 2, ignore_index=True)
        yb = np.concatenate([Z["y_ice_feb2025"].values[idx], Z["y_helene_sep2024"].values[idx]])
        Xb["s1"] = np.r_[np.ones(n), np.zeros(n)]
        Xb["s2"] = np.r_[np.zeros(n), np.ones(n)]
        m = PoissonRegressor(alpha=0.02, max_iter=3000, fit_intercept=False).fit(Xb.values, yb)
        cb = dict(zip(list(FEATS) + ["s1", "s2"], m.coef_))
        eb = sum(Xs[k].values * cb[k] for k in FEATS)
        eb = eb - np.log(np.exp(eb).sum()) + np.log(np.exp(e_final).sum())
        boot_eta.append(eb)
    boot_eta = np.array(boot_eta)
    eta_sd = boot_eta.std(axis=0)
    a_ice = severity_for_total(e_final, Z["y_ice_feb2025"].sum())
    a_hel = severity_for_total(e_final, Z["y_helene_sep2024"].sum())

    contrib = np.c_[[Xs[k].values * coef[k] for k in FEATS]].T

    loc_ice = dmg[(dmg.local) & (dmg.storm == "ice_feb2025")]
    n_oh_xf = int((xft.oh > 0).sum())
    n_ug_xf = int((xft.oh == 0).sum())
    loc_xf = xft.iloc[loc_ice.xi.values]
    r_oh = (loc_xf.oh > 0).sum() / max(n_oh_xf, 1)
    r_ug = (loc_xf.oh == 0).sum() / max(n_ug_xf, 1)
    log(f"local transformer/secondary damage rate in ice storm: OH {r_oh:.4f}, UG {r_ug:.4f}")

    zd = dmg[(dmg.storm != "other")]
    mix = zd.groupby("cls").size()
    mix = (mix / mix.sum()).round(4).to_dict()
    qty_dist = zd.qty.clip(upper=6).value_counts(normalize=True).sort_index()
    access_rate = float(zd.access.mean())
    special = zd.Special_Eq.replace("", "None").value_counts(normalize=True).round(4).to_dict()

    ice = dmg[(dmg.storm == "ice_feb2025") & dmg.rd.notna()].copy()
    start = pd.Timestamp("2025-02-12")
    ice["day_done"] = (ice.rd - start).dt.days
    obs_curve = [float((ice.day_done <= d).mean()) for d in range(0, 15)]
    n_ice_jobs = int((dmg.storm == "ice_feb2025").sum())

    labor = labor_model(os.path.join(raw, "Install_Labor_HackOHIO.xlsx"))

    log("writing outputs")
    zone_rows = []
    for i, d in enumerate(devs):
        info = dev_info[d]
        top = np.argsort(-contrib[i])[:3]
        zone_rows.append([
            info["kind"][0],
            di.get(parent.get(d), -1),
            int(Z.direct_cust[i]), int(Z.down_cust[i]),
            round(float(e_final[i]), 4), round(float(eta_sd[i]), 4),
            round(float(Z.pri_oh_ft[i] / 5280), 3), round(float(Z.pri_ug_ft[i] / 5280), 3),
            int(Z.poles[i]), round(float(Z.pole_age[i]), 1), round(float(Z.wood_frac[i]), 3),
            [int(t) for t in top], [round(float(contrib[i][t]), 3) for t in top],
            int(Z.crit[i] > 0),
            int(Z.y_ice_feb2025[i] + Z.y_helene_sep2024[i] + Z.y_other[i]),
            round(float(Z.oh_frac[i]), 3),
        ])
    feat_names = list(FEATS)

    cust = cust.reset_index(drop=True)
    used, pids = set(), []
    for pn in cust.premise_nu:
        k = 0
        h = hid(pn)
        while h in used:
            k += 1
            h = hid(pn, k)
        used.add(h)
        pids.append(h)
    cust["pid"] = pids
    cust["gx"] = (cust.lon / GRID).round().astype(int)
    cust["gy"] = (cust.lat / GRID).round().astype(int)
    cust["xi"] = cust.xf.map(xidx)
    prem = {
        "pid": cust.pid.tolist(),
        "xi": cust.xi.astype(int).tolist(),
        "gx": cust.gx.tolist(), "gy": cust.gy.tolist(),
        "svc": cust.kind_of_se.map({"Overhead": "O", "Underground": "U"}).fillna("?").tolist(),
        "city": cust.city.str.title().tolist(),
    }
    xfs = {
        "zi": xft.zi.astype(int).tolist(),
        "n": xft.n.astype(int).tolist(),
        "oh": (xft.oh > 0).astype(int).tolist(),
        "steps": xft.steps.tolist(),
    }
    cells = collections.Counter(zip(cust.gx, cust.gy))
    fac_pub = collections.Counter()
    for f in fac:
        fac_pub[(round(f["lon"] / 0.012), round(f["lat"] / 0.012), f["kind"])] += 1

    model = {
        "generated": dt.datetime.now().isoformat(timespec="seconds"),
        "excluded_features": {
            "Average pole age / share of pre-1975 poles": "No positive signal once line exposure and rural density are controlled for (urban poles are older but sheltered); excluded so explanations stay physically sensible.",
            "Overhead secondary miles": "Collinear with overhead primary; excluded.",
            "Life-support / THREAT flags": "Sensitive; never loaded by the pipeline.",
        },
        "features": feat_names, "friendly": FRIENDLY,
        "coef": {k: round(float(coef[k]), 4) for k in FEATS},
        "storm_intercepts": {"ice_feb2025": round(float(a_ice), 4), "helene_sep2024": round(float(a_hel), 4)},
        "observed_zone_damage": {s: int(Z[f"y_{s}"].sum()) for s in STORMS},
        "observed_records": {s: int((dmg.storm == s).sum()) for s in list(STORMS) + ["other"]},
        "local_rate": {"oh": round(float(r_oh), 5), "ug": round(float(r_ug), 5)},
        "backtests": backtests,
        "job_mix": mix, "qty_dist": {int(k): round(float(v), 4) for k, v in qty_dist.items()},
        "access_rate": round(access_rate, 4), "special_equipment": special,
        "obs_restoration_curve_ice": obs_curve, "n_ice_jobs": n_ice_jobs,
        "labor": labor,
        "grid": GRID, "k_anon": K_ANON,
        "data_notes": {
            "damage_records": int(len(dmg)), "premises": int(len(cust)), "zones": len(devs),
            "transformers": int(len(xft)), "circuits": int(cust.circ.nunique()),
            "median_snap_m": round(float(dmg.snap_m.median()), 1),
            "facilities_matched": len(fac),
        },
    }
    with open(os.path.join(OUT, "model.json"), "w") as f:
        json.dump(model, f)
    with open(os.path.join(OUT, "network.json"), "w") as f:
        json.dump({"zones": zone_rows, "xfs": xfs}, f, separators=(",", ":"))
    with open(os.path.join(OUT, "premises.json"), "w") as f:
        json.dump(prem, f, separators=(",", ":"))
    with open(os.path.join(OUT, "map.json"), "w") as f:
        json.dump({
            "cells": [[gx, gy, n] for (gx, gy), n in cells.items()],
            "backbone": backbone,
            "facilities": [[a, b, k, n] for (a, b, k), n in fac_pub.items()],
        }, f, separators=(",", ":"))
    log("done ->", os.path.abspath(OUT))
    for fn in os.listdir(OUT):
        log(f"  {fn}: {os.path.getsize(os.path.join(OUT, fn)) / 1e6:.2f} MB")


if __name__ == "__main__":
    main()
